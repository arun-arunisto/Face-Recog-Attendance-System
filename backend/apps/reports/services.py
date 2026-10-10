import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse

def build_attendance_csv_response(queryset, filename="attendance_report.csv"):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow([
        'Student ID', 'Student Name', 'Email', 'Department',
        'Class', 'Section', 'Date', 'Check-in Time', 'Check-out Time',
        'Status', 'Recognition Confidence', 'Camera / Source', 'Marked Manually'
    ])

    for record in queryset:
        writer.writerow([
            record.student.student_id,
            record.student.user.full_name,
            record.student.user.email,
            record.student.department.name,
            record.student.class_name,
            record.student.section,
            str(record.date),
            record.check_in_time.strftime("%H:%M:%S") if record.check_in_time else "",
            record.check_out_time.strftime("%H:%M:%S") if record.check_out_time else "",
            record.status,
            f"{record.recognition_confidence * 100:.1f}%",
            record.camera.name if record.camera else ("Manual" if record.marked_manually else "Webcam"),
            "Yes" if record.marked_manually else "No"
        ])

    return response


def build_attendance_xlsx_response(queryset, filename="attendance_report.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Records"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        'Student ID', 'Student Name', 'Email', 'Department',
        'Class', 'Section', 'Date', 'Check-in Time', 'Check-out Time',
        'Status', 'Confidence', 'Source', 'Manual'
    ]
    ws.append(headers)

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    # Rows styling
    status_fills = {
        'PRESENT': PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        'LATE': PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        'ABSENT': PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    }

    for row_idx, record in enumerate(queryset, start=2):
        row_values = [
            record.student.student_id,
            record.student.user.full_name,
            record.student.user.email,
            record.student.department.name,
            record.student.class_name,
            record.student.section,
            str(record.date),
            record.check_in_time.strftime("%H:%M:%S") if record.check_in_time else "--",
            record.check_out_time.strftime("%H:%M:%S") if record.check_out_time else "--",
            record.status,
            f"{record.recognition_confidence * 100:.1f}%",
            record.camera.name if record.camera else ("Manual" if record.marked_manually else "Webcam"),
            "Yes" if record.marked_manually else "No"
        ]
        ws.append(row_values)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 7, 8, 9, 10, 11, 13]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            # Highlight status cell
            if col_idx == 10 and record.status in status_fills:
                cell.fill = status_fills[record.status]

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
