import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS_DIR = os.path.join(BASE_DIR, 'diagrams')
SCREENSHOTS_DIR = os.path.join(BASE_DIR, 'screenshots')
OUTPUT_FILE = os.path.join(BASE_DIR, 'FINAL_PROJECT_REPORT.docx')

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: {"val": "single", "sz": "4", "color": "CCCCCC"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for border_name, border_props in kwargs.items():
        val = border_props.get('val', 'single')
        sz = border_props.get('sz', '4')
        color = border_props.get('color', 'CCCCCC')
        tag = f'<w:{border_name} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def add_page_number_to_footer(footer_para):
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_para.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 116, 139)

    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    footer_para._p.append(fldSimple)

    run2 = footer_para.add_run(" of ")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(100, 116, 139)

    fldSimple2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    footer_para._p.append(fldSimple2)

class ReportBuilder:
    def __init__(self):
        self.doc = Document()
        self.figure_count = 0
        self.table_count = 0
        self._configure_document()

    def _configure_document(self):
        # A4 Page Settings: 8.27 x 11.69 inches
        for section in self.doc.sections:
            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.0)

            # Header
            header = section.header
            header_para = header.paragraphs[0]
            header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hrun = header_para.add_run("Smart Face Recognition Attendance System | Final Year Project Report")
            hrun.font.name = "Times New Roman"
            hrun.font.size = Pt(8.5)
            hrun.font.color.rgb = RGBColor(148, 163, 184)

            # Footer
            footer = section.footer
            footer_para = footer.paragraphs[0]
            add_page_number_to_footer(footer_para)

        # Configure Normal Style
        normal = self.doc.styles['Normal']
        normal.font.name = 'Times New Roman'
        normal.font.size = Pt(12)
        normal.font.color.rgb = RGBColor(15, 23, 42)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.space_before = Pt(0)

    def add_chapter_heading(self, chapter_num, title):
        self.doc.add_page_break()
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        r_num = p.add_run(f"CHAPTER {chapter_num}\n")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(16)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(30, 41, 59)

        r_title = p.add_run(title.upper())
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(37, 99, 235)

    def add_section_heading(self, sec_num, title):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run = p.add_run(f"{sec_num} {title}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(13.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_subsection_heading(self, subsec_num, title):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run = p.add_run(f"{subsec_num} {title}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 65, 85)

    def add_paragraph(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        return p

    def add_callout(self, text, title="NOTE / ARCHITECTURAL SPECIFICATION", bg="#F8FAFC", border_color="#3B82F6"):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.0)
        set_cell_background(cell, bg.replace("#", ""))
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        set_cell_border(cell,
            left={"val": "single", "sz": "18", "color": border_color.replace("#", "")},
            top={"val": "none", "sz": "0", "color": "FFFFFF"},
            bottom={"val": "none", "sz": "0", "color": "FFFFFF"},
            right={"val": "none", "sz": "0", "color": "FFFFFF"}
        )

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"[{title}]\n")
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(37, 99, 235)

        r_text = p.add_run(text)
        r_text.font.name = "Times New Roman"
        r_text.font.size = Pt(10.5)
        r_text.font.italic = True
        r_text.font.color.rgb = RGBColor(51, 65, 85)

        # Space after callout
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(4)

    def add_code_block(self, code_text, caption=""):
        if caption:
            p_cap = self.doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(2)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(f"Listing: {caption}")
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(10)
            r_cap.font.bold = True
            r_cap.font.color.rgb = RGBColor(71, 85, 105)

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.0)
        set_cell_background(cell, "0F172A") # Dark slate
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        set_cell_border(cell,
            left={"val": "single", "sz": "6", "color": "334155"},
            right={"val": "single", "sz": "6", "color": "334155"},
            top={"val": "single", "sz": "6", "color": "334155"},
            bottom={"val": "single", "sz": "6", "color": "334155"}
        )

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(code_text.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(241, 245, 249)

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(6)

    def add_figure(self, img_path, caption, width=Inches(5.8)):
        self.figure_count += 1
        if not os.path.exists(img_path):
            print(f"Warning: Figure image not found: {img_path}")
            return

        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)

        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    def add_table_data(self, headers, rows, caption, col_widths=None):
        self.table_count += 1
        # Table Caption Above Table
        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(30, 41, 59)

        table = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            cell = hdr_cells[i]
            if col_widths and i < len(col_widths):
                cell.width = Inches(col_widths[i])
            set_cell_background(cell, "1E293B") # Dark slate
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_border(cell,
                bottom={"val": "single", "sz": "8", "color": "0F172A"},
                top={"val": "single", "sz": "4", "color": "CBD5E1"},
                left={"val": "single", "sz": "4", "color": "CBD5E1"},
                right={"val": "single", "sz": "4", "color": "CBD5E1"}
            )
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(header_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                cell = row_cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = Inches(col_widths[c_idx])
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                set_cell_border(cell,
                    bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                    top={"val": "single", "sz": "4", "color": "E2E8F0"},
                    left={"val": "single", "sz": "4", "color": "E2E8F0"},
                    right={"val": "single", "sz": "4", "color": "E2E8F0"}
                )
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(cell_value))
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)

        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(6)

    def save(self):
        print(f"Saving final report to {OUTPUT_FILE}...")
        self.doc.save(OUTPUT_FILE)
        print(f"Report saved successfully! Size: {os.path.getsize(OUTPUT_FILE)} bytes")

print("ReportBuilder framework initialized successfully.")
